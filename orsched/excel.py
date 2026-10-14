# -*- coding: utf-8 -*-
"""Writing several DataFrames to one formatted Excel workbook.

"""
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import pandas as pd


def write_workbook(path, sheets, percent_columns=()):
    """sheets: {sheet name: DataFrame}. Bold header, frozen first row,
    fitted column widths; columns in percent_columns get a % format."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            name = name[:31]  # Excel limit
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            for i, col in enumerate(df.columns, start=1):
                ws.cell(row=1, column=i).font = Font(bold=True)
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)])
                ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 40)
                if col in percent_columns:
                    for r in range(2, len(df) + 2):
                        ws.cell(row=r, column=i).number_format = "0.00%"
